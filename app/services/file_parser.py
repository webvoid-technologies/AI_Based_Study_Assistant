"""Extract plain text from every supported study-material format."""

import json
import os
import shutil
from pathlib import Path

import fitz
import pytesseract
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageOps
from pptx import Presentation


class FileParsingError(ValueError):
    """Raised when an uploaded file cannot be converted to text."""


class FileParser:
    supported_extensions = {".pdf", ".docx", ".pptx", ".xlsx", ".png", ".jpg", ".jpeg", ".json", ".txt"}

    def extract_text(self, file_path: Path) -> str:
        extension = file_path.suffix.lower()
        if extension not in self.supported_extensions:
            raise FileParsingError("Unsupported file format.")

        parsers = {
            ".pdf": self._pdf,
            ".docx": self._docx,
            ".pptx": self._pptx,
            ".xlsx": self._xlsx,
            ".png": self._image,
            ".jpg": self._image,
            ".jpeg": self._image,
            ".json": self._json,
            ".txt": self._text,
        }
        try:
            content = parsers[extension](file_path).strip()
        except FileParsingError:
            raise
        except Exception as exc:
            if extension in {".png", ".jpg", ".jpeg"}:
                raise FileParsingError("OCR failed. Check that Tesseract is installed and the image is readable.") from exc
            raise FileParsingError("The file could not be read. It may be damaged or invalid.") from exc

        if not content:
            raise FileParsingError("No readable text was found in the document.")
        return content

    @staticmethod
    def _pdf(file_path: Path) -> str:
        with fitz.open(file_path) as document:
            return "\n\n".join(page.get_text("text") for page in document)

    @staticmethod
    def _docx(file_path: Path) -> str:
        document = Document(file_path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join(paragraphs + rows)

    @staticmethod
    def _pptx(file_path: Path) -> str:
        presentation = Presentation(file_path)
        slides: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            text = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if text:
                slides.append(f"Slide {index}\n" + "\n".join(text))
        return "\n\n".join(slides)

    @staticmethod
    def _xlsx(file_path: Path) -> str:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        output: list[str] = []
        try:
            for sheet in workbook.worksheets:
                output.append(f"Sheet: {sheet.title}")
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                headers = [str(value).strip() if value is not None else f"Column {index + 1}" for index, value in enumerate(rows[0])]
                for row in rows[1:]:
                    values = [f"{headers[index]} {value}" for index, value in enumerate(row) if value not in (None, "")]
                    if values:
                        output.append(" ".join(values))
        finally:
            workbook.close()
        return "\n".join(output)

    @staticmethod
    def _image(file_path: Path) -> str:
        command = os.getenv("TESSERACT_CMD") or shutil.which("tesseract")
        if command is None:
            for candidate in ("/opt/homebrew/bin/tesseract", "/usr/local/bin/tesseract", "/usr/bin/tesseract"):
                if Path(candidate).exists():
                    command = candidate
                    break
        if command is None:
            raise FileParsingError("Tesseract OCR is not installed or is unavailable.")
        pytesseract.pytesseract.tesseract_cmd = command

        with Image.open(file_path) as image:
            processed = ImageOps.autocontrast(image.convert("L"))
            return pytesseract.image_to_string(processed)

    @staticmethod
    def _json(file_path: Path) -> str:
        with file_path.open(encoding="utf-8") as stream:
            return json.dumps(json.load(stream), indent=2, ensure_ascii=False)

    @staticmethod
    def _text(file_path: Path) -> str:
        return file_path.read_text(encoding="utf-8-sig", errors="replace")


file_parser = FileParser()
